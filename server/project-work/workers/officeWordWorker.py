#!/usr/bin/env python3
"""Fixed, stdin-only DOCX builder used by Pi Agent's Office artifact service."""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


MAX_STDIN_BYTES = 1_500_000
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def fail(message: str) -> None:
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
    raise SystemExit(1)


def read_payload() -> dict:
    raw = sys.stdin.buffer.read(MAX_STDIN_BYTES + 1)
    if not raw or len(raw) > MAX_STDIN_BYTES:
        fail("invalid request size")
    try:
        payload = json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        fail("invalid request JSON")
    if not isinstance(payload, dict) or set(payload) != {"request"}:
        fail("invalid request envelope")
    request = payload["request"]
    if not isinstance(request, dict):
        fail("invalid request")
    return request


def contains_cjk(text: str) -> bool:
    return any(
        "\u3400" <= character <= "\u9fff"
        or "\uf900" <= character <= "\ufaff"
        for character in text
    )


def set_run_font(run, latin: str = "Calibri", east_asia: str = "Hiragino Sans GB") -> None:
    if contains_cjk(run.text or ""):
        latin = east_asia
    run.font.name = latin
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:cs"), east_asia)
    language = run_properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        run_properties.append(language)
    language.set(qn("w:eastAsia"), "zh-CN")


def set_style_font(style, latin: str, size: float, color: str | None = None) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    style_properties = style.element.get_or_add_rPr()
    fonts = style_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        style_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), "PingFang SC")


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    set_style_font(normal, "Calibri", 11, "1F2933")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 16, 8),
        "Heading 2": (13, "2E74B5", 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        set_style_font(style, "Calibri", size, color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        set_style_font(style, "Calibri", 11, "1F2933")
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    title = styles.add_style("Pi Title", 1)
    set_style_font(title, "Calibri", 24, "0B2545")
    title.font.bold = True
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True

    subtitle = styles.add_style("Pi Subtitle", 1)
    set_style_font(subtitle, "Calibri", 11, "5F6B76")
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.keep_with_next = True


def configure_page(document: Document) -> None:
    for section in document.sections:
        section.start_type = WD_SECTION.NEW_PAGE
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("6B7280")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_furniture(document: Document, title: str) -> None:
    for section in document.sections:
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(title[:90])
        set_run_font(run)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("6B7280")
        add_page_field(section.footer.paragraphs[0])


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_properties.append(header)


def set_row_no_split(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    row_properties.append(cannot_split)


def set_cell_margins(cell) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for side, width in CELL_MARGINS_DXA.items():
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    cell.width = Twips(width)
    cell_properties = cell._tc.get_or_add_tcPr()
    cell_width = cell_properties.get_or_add_tcW()
    cell_width.set(qn("w:w"), str(width))
    cell_width.set(qn("w:type"), "dxa")


def set_cell_fill(cell, color: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), color)
    shading.set(qn("w:val"), "clear")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table_properties = table._tbl.tblPr

    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.insert(0, table_width)
    table_width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")

    table_indent = table_properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    table_indent.set(qn("w:type"), "dxa")

    layout = table_properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    borders = table_properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "D8DEE6")

    for row in table.rows:
        set_row_no_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def column_widths(headers: list[str], rows: list[list[str]]) -> list[int]:
    weights = []
    for column, header in enumerate(headers):
        samples = [header] + [row[column] for row in rows[:30]]
        longest = max(4, min(60, max(len(str(value)) for value in samples)))
        weights.append(max(8, longest))
    total = sum(weights)
    widths = [max(720, round(CONTENT_WIDTH_DXA * weight / total)) for weight in weights]
    difference = CONTENT_WIDTH_DXA - sum(widths)
    widths[-1] += difference
    if widths[-1] < 720:
        deficit = 720 - widths[-1]
        donor = max(range(len(widths) - 1), key=lambda index: widths[index])
        widths[donor] -= deficit
        widths[-1] = 720
    return widths


def style_table_text(cell, *, header: bool = False) -> None:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            set_run_font(run)
            run.font.size = Pt(10)
            run.font.bold = header
            run.font.color.rgb = RGBColor.from_string("1F2933")


def add_table(document: Document, table_request: dict) -> None:
    headers = table_request["headers"]
    rows = table_request["rows"]
    widths = column_widths(headers, rows)
    table = document.add_table(rows=1, cols=len(headers))
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, value in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = value
        set_cell_fill(cell, "F2F4F7")
        style_table_text(cell, header=True)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            style_table_text(cells[index])
    set_table_geometry(table, widths)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def build_document(request: dict) -> dict:
    expected = {"fileName", "title", "subtitle", "sections"}
    if set(request) != expected:
        fail("request was not normalized")
    file_name = request["fileName"]
    if Path(file_name).name != file_name or not file_name.lower().endswith(".docx"):
        fail("invalid output file name")

    document = Document()
    configure_page(document)
    configure_styles(document)
    configure_furniture(document, request["title"])
    document.core_properties.title = request["title"]
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""

    title = document.add_paragraph(style="Pi Title")
    title_run = title.add_run(request["title"])
    set_run_font(title_run)
    if request["subtitle"]:
        subtitle = document.add_paragraph(style="Pi Subtitle")
        subtitle_run = subtitle.add_run(request["subtitle"])
        set_run_font(subtitle_run)

    paragraph_count = 0
    list_item_count = 0
    table_count = 0
    table_cell_count = 0
    for section in request["sections"]:
        if section["heading"]:
            paragraph = document.add_heading(section["heading"], level=section["level"])
            for run in paragraph.runs:
                set_run_font(run)
            paragraph_count += 1
        for text in section["paragraphs"]:
            paragraph = document.add_paragraph(text)
            for run in paragraph.runs:
                set_run_font(run)
            paragraph_count += 1
        for text in section["bullets"]:
            paragraph = document.add_paragraph(text, style="List Bullet")
            for run in paragraph.runs:
                set_run_font(run)
            paragraph_count += 1
            list_item_count += 1
        for text in section["numbered"]:
            paragraph = document.add_paragraph(text, style="List Number")
            for run in paragraph.runs:
                set_run_font(run)
            paragraph_count += 1
            list_item_count += 1
        for table_request in section["tables"]:
            add_table(document, table_request)
            table_count += 1
            table_cell_count += len(table_request["headers"]) * (1 + len(table_request["rows"]))

    output_path = Path.cwd() / file_name
    document.save(output_path)
    os.chmod(output_path, 0o600)

    round_trip = Document(output_path)
    if not round_trip.paragraphs or not round_trip.styles["Pi Title"]:
        fail("DOCX structural round-trip failed")
    return {
        "ok": True,
        "fileName": file_name,
        "structure": {
            "preset": "standard_business_brief",
            "pageSize": "letter",
            "paragraphCount": paragraph_count,
            "listItemCount": list_item_count,
            "tableCount": table_count,
            "tableCellCount": table_cell_count,
            "sectionCount": len(request["sections"]),
        },
    }


def main() -> None:
    request = read_payload()
    try:
        result = build_document(request)
    except SystemExit:
        raise
    except Exception as error:  # The parent process deliberately receives no traceback.
        fail(f"DOCX build failed: {type(error).__name__}")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
